"""Essay Scorer — reads one essay (PDF or DOCX) and scores it 1-5 on
Structure, Evidence, and Vocabulary Range, out of 15 total.

Deterministic: fixed rules and regex only, no AI, no randomness.
Thresholds are the ones locked in scoped-problem.docx.
"""

import re
import sys
from pathlib import Path

import docx
from pypdf import PdfReader

PHRASES = ["for example", "for instance", "such as", "according to"]
MIN_WORDS_FOR_VOCAB = 50


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        document = docx.Document(str(path))
        paragraphs = [
            p.text.strip()
            for p in document.paragraphs
            if p.text.strip() and not p.style.name.startswith(("Heading", "Title"))
        ]
        return "\n\n".join(paragraphs)
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(pages)
    raise ValueError(f"Unsupported file type: {suffix} (use .pdf or .docx)")


def split_paragraphs(text: str):
    return [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]


def split_sentences(text: str):
    text = text.replace("\n", " ")
    parts = re.split(r'(?:(?<=[.!?])|(?<=[.!?]["\'”’)]))\s+', text)
    return [s.strip() for s in parts if s.strip()]


def find_evidence_markers(text: str):
    markers = []
    for match in re.finditer(r'"([^"]{3,})"', text):
        markers.append(match.group(0))
    for match in re.finditer(r"“([^”]{3,})”", text):
        markers.append(match.group(0))
    for phrase in PHRASES:
        for match in re.finditer(re.escape(phrase), text, re.IGNORECASE):
            markers.append(match.group(0))
    return markers


def score_structure(paragraphs):
    count = len(paragraphs)
    if count >= 5:
        score = 5
    elif count == 4:
        score = 4
    elif count == 3:
        score = 3
    elif count == 2:
        score = 2
    else:
        score = 1

    if count == 0:
        evidence = "not enough text"
    else:
        opening = " ".join(paragraphs[0].split()[:12])
        ellipsis = "..." if len(paragraphs[0].split()) > 12 else ""
        evidence = f"{count} paragraph(s) detected; opens with: \"{opening}{ellipsis}\""
    return score, evidence


def score_evidence(text: str):
    markers = find_evidence_markers(text)
    count = len(markers)
    if count >= 4:
        score = 5
    elif count == 3:
        score = 4
    elif count == 2:
        score = 3
    elif count == 1:
        score = 2
    else:
        score = 1

    if count == 0:
        return score, "not enough text"

    marker = markers[0]
    for sentence in split_sentences(text):
        if marker.strip('"“”').lower() in sentence.lower():
            return score, sentence
    return score, marker


def score_vocabulary(text: str):
    words = re.findall(r"[A-Za-z']+", text)
    total_words = len(words)
    if total_words < MIN_WORDS_FOR_VOCAB:
        return 1, f"not enough text ({total_words} words)"

    long_words = [w.lower() for w in words if len(w) >= 4]
    unique_long = set(long_words)
    ratio = len(unique_long) / len(long_words) if long_words else 0.0

    if ratio >= 0.60:
        score = 5
    elif ratio >= 0.50:
        score = 4
    elif ratio >= 0.40:
        score = 3
    elif ratio >= 0.30:
        score = 2
    else:
        score = 1

    evidence = (
        f"{len(unique_long)} unique 4+ letter words out of {len(long_words)} "
        f"(ratio {ratio:.2f}) across {total_words} total words"
    )
    return score, evidence


def evaluate_essay(path: Path) -> dict:
    text = extract_text(path)
    paragraphs = split_paragraphs(text)

    structure_score, structure_evidence = score_structure(paragraphs)
    evidence_score, evidence_evidence = score_evidence(text)
    vocab_score, vocab_evidence = score_vocabulary(text)

    total = structure_score + evidence_score + vocab_score

    return {
        "file": str(path),
        "structure": {"score": structure_score, "evidence": structure_evidence},
        "evidence": {"score": evidence_score, "evidence": evidence_evidence},
        "vocabulary_range": {"score": vocab_score, "evidence": vocab_evidence},
        "total": total,
    }


def print_results(results: dict) -> None:
    print(f"File: {results['file']}")
    print(f"Structure: {results['structure']['score']}/5 -- {results['structure']['evidence']}")
    print(f"Evidence: {results['evidence']['score']}/5 -- {results['evidence']['evidence']}")
    print(
        f"Vocabulary Range: {results['vocabulary_range']['score']}/5 -- "
        f"{results['vocabulary_range']['evidence']}"
    )
    print(f"Total: {results['total']}/15")


def write_report(path: Path, results: dict) -> Path:
    report = docx.Document()
    report.add_heading("Essay Scorer Report", level=1)
    report.add_paragraph(f"File: {results['file']}")
    report.add_paragraph(
        f"Structure: {results['structure']['score']}/5 -- {results['structure']['evidence']}"
    )
    report.add_paragraph(
        f"Evidence: {results['evidence']['score']}/5 -- {results['evidence']['evidence']}"
    )
    report.add_paragraph(
        f"Vocabulary Range: {results['vocabulary_range']['score']}/5 -- "
        f"{results['vocabulary_range']['evidence']}"
    )
    report.add_paragraph(f"Total: {results['total']}/15")

    report_path = path.with_name(f"{path.stem}-report.docx")
    report.save(str(report_path))
    return report_path


def main():
    if len(sys.argv) != 2:
        print("Usage: python evaluate.py <file.pdf or .docx>")
        sys.exit(1)

    path = Path(sys.argv[1])
    results = evaluate_essay(path)
    print_results(results)
    report_path = write_report(path, results)
    print(f"Report written to: {report_path}")


if __name__ == "__main__":
    main()
